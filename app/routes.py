# This file defines all the routes for the Flask application.
# It handles authentication, dashboard, beneficiary management (CRUD),
# approval workflows, reports, data import/export, and user settings.

import json
from datetime import datetime
from io import BytesIO
from functools import wraps

from flask import (render_template, request, redirect, url_for, flash, jsonify,
                   send_file, make_response, Blueprint)
from flask_login import (login_user, logout_user, login_required, current_user)
from sqlalchemy import func, extract, or_
import openpyxl
from openpyxl import Workbook
from dateutil.relativedelta import relativedelta

from . import db
from .models import User, Record, PendingChange
from .utils import calculate_age_group, generate_beneficiary_pdf

# Create a Blueprint to organize routes. All routes are attached to this blueprint.
bp = Blueprint('routes', __name__, template_folder='templates')


# --- Custom Decorators ---
def admin_required(f):
    """
    A decorator to restrict access to a route to admin users only.
    If the user is not an admin, they are redirected to the dashboard.
    """
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin():
            flash('Admin access is required for this page.', 'danger')
            return redirect(url_for('routes.dashboard'))
        return f(*args, **kwargs)
    return decorated_function


# --- Authentication Routes ---
@bp.route('/')
def index():
    """
    Root URL route. Redirects to the login page if not authenticated,
    otherwise redirects to the dashboard.
    """
    if not current_user.is_authenticated:
        return redirect(url_for('routes.login'))
    return redirect(url_for('routes.dashboard'))

@bp.route('/login', methods=['GET', 'POST'])
def login():
    """
    Handles user login. Displays the login form and processes credentials.
    """
    if current_user.is_authenticated:
        return redirect(url_for('routes.dashboard'))

    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            login_user(user, remember=True)
            flash('تم تسجيل الدخول بنجاح', 'success')
            return redirect(url_for('routes.dashboard'))
        else:
            flash('اسم المستخدم أو كلمة المرور غير صحيحة', 'error')
    return render_template('login.html')

@bp.route('/logout')
@login_required
def logout():
    """
    Handles user logout.
    """
    logout_user()
    flash('تم تسجيل الخروج بنجاح', 'info')
    return redirect(url_for('routes.login'))


# --- Main Application Routes ---
@bp.route('/dashboard')
@login_required
def dashboard():
    """
    Displays the main dashboard with statistics and recent activity.
    """
    # --- Statistics Queries ---
    total_beneficiaries = Record.query.filter_by(is_deleted=False).count()
    current_month = datetime.now().month
    current_year = datetime.now().year
    new_this_month = Record.query.filter_by(is_deleted=False).filter(
        extract('month', Record.created_at) == current_month,
        extract('year', Record.created_at) == current_year
    ).count()
    needs_review = Record.query.filter_by(status='بحاجة لمراجعة', is_deleted=False).count()
    inactive_records = Record.query.filter_by(status='غير نشط', is_deleted=False).count()

    # --- Recent Activity Query ---
    recent_beneficiaries = Record.query.filter_by(is_deleted=False).order_by(Record.created_at.desc()).limit(5).all()

    # --- Pending Changes Count for Admins ---
    pending_changes_count = 0
    if current_user.is_admin():
        pending_changes_count = PendingChange.query.filter_by(status='pending').count()

    stats = {
        'total_beneficiaries': total_beneficiaries, 'new_this_month': new_this_month,
        'needs_review': needs_review, 'inactive_records': inactive_records,
        'pending_changes_count': pending_changes_count
    }
    return render_template('dashboard.html', stats=stats, recent_beneficiaries=recent_beneficiaries)


# --- Beneficiary CRUD Routes ---
@bp.route('/beneficiaries')
@login_required
def beneficiaries():
    """
    Displays the main beneficiary management page with filtering, search, and pagination.
    Only heads of household are displayed at the top level. Family members are nested.
    """
    page = request.args.get('page', 1, type=int)
    per_page = 10  # This will apply to heads of household

    # Base query: only select records that are not soft-deleted and are heads of household.
    query = Record.query.filter_by(is_deleted=False).filter(Record.head_of_household_id.is_(None))

    # --- Search and Filtering Logic ---
    search = request.args.get('search', '')
    if search:
        search_term = f"%{search}%"
        # A simple search on heads of household.
        # A more complex search would require joining or subqueries to include family members.
        query = query.filter(or_(
            Record.first_name.ilike(search_term), Record.father_name.ilike(search_term),
            Record.grandfather_name.ilike(search_term), Record.family_name.ilike(search_term),
            Record.id_passport_number.ilike(search_term), Record.phone_number.ilike(search_term)
        ))
    status_filter = request.args.get('status', '')
    if status_filter:
        query = query.filter_by(status=status_filter)
    start_date_str = request.args.get('start_date', '')
    if start_date_str:
        query = query.filter(Record.created_at >= datetime.strptime(start_date_str, '%Y-%m-%d').date())
    end_date_str = request.args.get('end_date', '')
    if end_date_str:
        query = query.filter(Record.created_at <= datetime.strptime(end_date_str, '%Y-%m-%d').date())

    # Paginate the results
    records = query.order_by(Record.created_at.desc()).paginate(page=page, per_page=per_page, error_out=False)

    # Fetch potential heads of household for the 'Add/Edit' modals.
    potential_heads = Record.query.filter(Record.head_of_household_id.is_(None)).order_by(Record.first_name, Record.family_name).all()

    return render_template('beneficiaries.html', records=records, search=search,
                           status_filter=status_filter, start_date=start_date_str, end_date=end_date_str,
                           potential_heads_of_household=potential_heads)

@bp.route('/add_single_beneficiary', methods=['GET', 'POST'])
@login_required
def add_single_beneficiary():
    """
    Handles the creation of a new beneficiary record from a dedicated page.
    """
    if request.method == 'POST':
        try:
            head_of_household_id_str = request.form.get('head_of_household_id')
            head_of_household_id = int(head_of_household_id_str) if head_of_household_id_str else None

            # Basic validation to prevent a record from being its own head of household.
            if head_of_household_id and 'id' in request.form and int(request.form['id']) == head_of_household_id:
                flash('لا يمكن تعيين المستفيد كرب أسرة لنفسه.', 'error')
                potential_heads = Record.query.filter(Record.head_of_household_id.is_(None)).order_by(Record.first_name, Record.family_name).all()
                return render_template('add_single_beneficiary.html', potential_heads_of_household=potential_heads, existing_record=request.form), 400

            # Check for duplicate id_passport_number
            existing_record = Record.query.filter_by(id_passport_number=request.form['id_passport_number']).first()
            if existing_record:
                flash('رقم الهوية/جواز السفر موجود بالفعل.', 'error')
                potential_heads = Record.query.filter(Record.head_of_household_id.is_(None)).order_by(Record.first_name, Record.family_name).all()
                return render_template('add_single_beneficiary.html', potential_heads_of_household=potential_heads, existing_record=request.form), 400

            record = Record(
                first_name=request.form['first_name'],
                father_name=request.form['father_name'],
                grandfather_name=request.form['grandfather_name'],
                family_name=request.form['family_name'],
                id_passport_number=request.form['id_passport_number'],
                date_of_birth=datetime.strptime(request.form['date_of_birth'], '%Y-%m-%d').date(),
                gender=request.form['gender'],
                marital_status=request.form['marital_status'],
                phone_number=request.form.get('phone_number', ''),
                address=request.form.get('address', ''),
                status=request.form['status'],
                created_by_user_id=current_user.id,
                head_of_household_id=head_of_household_id
            )
            db.session.add(record)
            db.session.commit()
            flash('تم إضافة المستفيد بنجاح', 'success')
            return redirect(url_for('routes.beneficiaries'))
        except Exception as e:
            db.session.rollback()
            flash(f'حدث خطأ أثناء إضافة المستفيد: {str(e)}', 'error')
            potential_heads = Record.query.filter(Record.head_of_household_id.is_(None)).order_by(Record.first_name, Record.family_name).all()
            return render_template('add_single_beneficiary.html', potential_heads_of_household=potential_heads, existing_record=request.form), 500

    # For GET requests, render the form with potential heads of household.
    potential_heads = Record.query.filter(Record.head_of_household_id.is_(None)).order_by(Record.first_name, Record.family_name).all()
    return render_template('add_single_beneficiary.html', potential_heads_of_household=potential_heads)

def _get_updated_data_from_form(form):
    """Extracts updated data from the form."""
    return {
        'first_name': form['first_name'], 'father_name': form['father_name'],
        'grandfather_name': form['grandfather_name'], 'family_name': form['family_name'],
        'id_passport_number': form['id_passport_number'],
        'date_of_birth': form['date_of_birth'],
        'gender': form['gender'], 'marital_status': form['marital_status'],
        'phone_number': form.get('phone_number', ''),
        'address': form.get('address', ''), 'status': form['status'],
        'head_of_household_id': int(form.get('head_of_household_id')) if form.get('head_of_household_id') else None,
    }

def _update_record_directly(record, updated_data):
    """Updates the record directly."""
    for key, value in updated_data.items():
        if key == 'date_of_birth': value = datetime.strptime(value, '%Y-%m-%d').date()
        setattr(record, key, value)
    record.updated_at = datetime.utcnow()
    db.session.commit()
    flash('تم تحديث بيانات المستفيد بنجاح', 'success')

def _create_pending_change(record, updated_data):
    """Creates a pending change request for a non-admin user."""
    original_data = {
        'first_name': record.first_name, 'father_name': record.father_name,
        'grandfather_name': record.grandfather_name, 'family_name': record.family_name,
        'id_passport_number': record.id_passport_number,
        'date_of_birth': record.date_of_birth.strftime('%Y-%m-%d'),
        'gender': record.gender, 'marital_status': record.marital_status,
        'phone_number': record.phone_number or '',
        'address': record.address or '', 'status': record.status,
        'head_of_household_id': record.head_of_household_id
    }

    diff_data = {k: updated_data[k] for k, v in updated_data.items() if updated_data[k] != original_data.get(k)}

    if not diff_data:
        flash('لم يتم العثور على أي تغييرات لطلب التحديث.', 'info')
        return

    pending_change = PendingChange(
        record_id=record.id, user_id=current_user.id, change_type='update',
        changed_data=json.dumps(diff_data, default=str), status='pending'
    )
    db.session.add(pending_change)
    db.session.commit()
    flash('تم إرسال طلب التحديث للمراجعة', 'info')

@bp.route('/edit_beneficiary/<int:record_id>', methods=['POST'])
@login_required
def edit_beneficiary(record_id):
    """
    Handles editing an existing beneficiary record.
    Admins can edit directly. Regular users submit changes for approval.
    """
    record = Record.query.get_or_404(record_id)
    try:
        updated_data = _get_updated_data_from_form(request.form)

        if updated_data['head_of_household_id'] == record_id:
            flash('لا يمكن تعيين المستفيد كرب أسرة لنفسه.', 'error')
            return redirect(url_for('routes.beneficiaries'))

        if current_user.is_admin():
            _update_record_directly(record, updated_data)
        else:
            _create_pending_change(record, updated_data)
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء تحديث البيانات: {str(e)}', 'error')
    return redirect(url_for('routes.beneficiaries'))

@bp.route('/delete_beneficiary/<int:record_id>', methods=['POST'])
@login_required
def delete_beneficiary(record_id):
    """
    Handles soft deletion of a beneficiary record.
    Admins can soft delete directly. Regular users submit changes for approval.
    Prevents deletion of a head of household with dependents.
    """
    record_to_delete = Record.query.get_or_404(record_id)
    try:
        if current_user.is_admin():
            # Prevent deleting a head of household who has active family members.
            if record_to_delete.head_of_household_id is None and record_to_delete.family_members.filter_by(is_deleted=False).count() > 0:
                flash('لا يمكن حذف رب أسرة لديه أفراد أسرة مرتبطين به. يرجى أولاً إزالة أفراد الأسرة أو إعادة تعيينهم.', 'error')
                return redirect(url_for('routes.beneficiaries'))

            record_to_delete.is_deleted = True
            db.session.commit()
            flash('تم حذف المستفيد بنجاح.', 'success')
        else:
            # For non-admins, create a pending change request for soft deletion.
            existing_pending_delete = PendingChange.query.filter_by(
                record_id=record_id,
                change_type='delete',
                status='pending'
            ).first()
            if existing_pending_delete:
                flash('يوجد طلب حذف معلق بالفعل لهذا المستفيد.', 'info')
            else:
                pending_change = PendingChange(
                    record_id=record_id, user_id=current_user.id,
                    change_type='delete', status='pending'
                )
                db.session.add(pending_change)
                db.session.commit()
                flash('تم إرسال طلب الحذف للمراجعة', 'info')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء الحذف: {str(e)}', 'error')
    return redirect(url_for('routes.beneficiaries'))


# --- Approval Workflow Routes ---
@bp.route('/review_changes')
@admin_required
def review_changes():
    """
    Displays the page for admins to review pending change requests.
    """
    pending_changes = PendingChange.query.filter_by(status='pending').order_by(PendingChange.created_at.desc()).all()
    return render_template('review_changes.html', pending_changes=pending_changes)

@bp.route('/approve_change/<int:change_id>')
@admin_required
def approve_change(change_id):
    """
    Handles the approval of a pending change request by an admin.
    """
    change = PendingChange.query.get_or_404(change_id)
    try:
        if change.change_type == 'update':
            # Apply the updated data to the record.
            record = Record.query.get(change.record_id)
            if record:
                updated_data = json.loads(change.changed_data)
                for key, value in updated_data.items():
                    if key == 'date_of_birth':
                        value = datetime.strptime(value, '%Y-%m-%d').date()
                    elif key == 'head_of_household_id':
                        value = int(value) if value is not None and str(value).lower() not in ["", "none", "null"] else None
                    setattr(record, key, value)
                record.updated_at = datetime.utcnow()
        elif change.change_type == 'delete':
            # Soft delete the record, but check for dependents first.
            record_to_delete = Record.query.get(change.record_id)
            if record_to_delete:
                if record_to_delete.head_of_household_id is None and record_to_delete.family_members.filter_by(is_deleted=False).count() > 0:
                    flash('لا يمكن حذف رب أسرة لديه أفراد أسرة مرتبطين به. تم رفض الطلب تلقائياً.', 'error')
                    change.status = 'rejected'
                    change.reviewed_by = current_user.id
                    change.reviewed_at = datetime.utcnow()
                    db.session.commit()
                    return redirect(url_for('routes.review_changes'))

                record_to_delete.is_deleted = True

        # Mark the change as approved.
        change.status = 'approved'
        change.reviewed_by = current_user.id
        change.reviewed_at = datetime.utcnow()
        db.session.commit()
        flash('تم الموافقة على التغيير بنجاح', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء الموافقة: {str(e)}', 'error')
    return redirect(url_for('routes.review_changes'))

@bp.route('/reject_change/<int:change_id>')
@admin_required
def reject_change(change_id):
    """
    Handles the rejection of a pending change request by an admin.
    """
    change = PendingChange.query.get_or_404(change_id)
    try:
        change.status = 'rejected'
        change.reviewed_by = current_user.id
        change.reviewed_at = datetime.utcnow()
        db.session.commit()
        flash('تم رفض التغيير', 'info')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء الرفض: {str(e)}', 'error')
    return redirect(url_for('routes.review_changes'))


# --- Reports and Data Handling ---
@bp.route('/reports')
@login_required
def reports():
    """
    Generates and displays various reports and charts about the beneficiaries.
    """
    # --- Data Queries for Charts ---
    gender_data = db.session.query(Record.gender, func.count(Record.id)).filter(Record.is_deleted == False).group_by(Record.gender).all()
    marital_data = db.session.query(Record.marital_status, func.count(Record.id)).filter(Record.is_deleted == False).group_by(Record.marital_status).all()
    status_data = db.session.query(Record.status, func.count(Record.id)).filter(Record.is_deleted == False).group_by(Record.status).all()

    # Age group calculation
    records = Record.query.filter_by(is_deleted=False).all()
    age_groups = {'0-18': 0, '19-35': 0, '36-50': 0, '51-65': 0, '65+': 0}
    for record in records:
        age_group = calculate_age_group(record.age)
        age_groups[age_group] += 1

    # Monthly growth calculation for the last 12 months
    monthly_data = []
    for i in range(12):
        month_date = (datetime.now().replace(day=1) - relativedelta(months=i))
        count = Record.query.filter(extract('month', Record.created_at) == month_date.month,
                                    extract('year', Record.created_at) == month_date.year).count()
        monthly_data.append({'label': month_date.strftime('%Y-%m'), 'value': count})

    reports_data = {
        'gender': [{'label': item[0], 'value': item[1]} for item in gender_data],
        'marital_status': [{'label': item[0], 'value': item[1]} for item in marital_data],
        'status': [{'label': item[0], 'value': item[1]} for item in status_data],
        'age_groups': [{'label': k, 'value': v} for k, v in age_groups.items()],
        'monthly_growth': list(reversed(monthly_data))
    }
    return render_template('reports.html', reports_data=reports_data)

@bp.route('/import_export')
@login_required
def import_export():
    """
    Displays the page for importing and exporting data.
    """
    return render_template('import_export.html')

@bp.route('/download_template')
@login_required
def download_template():
    """
    Generates and serves an Excel template for importing beneficiaries.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Beneficiaries Template"
    headers = ['الاسم الأول', 'اسم الأب', 'اسم الجد', 'اسم العائلة', 'رقم الهوية/جواز السفر',
               'تاريخ الميلاد', 'الجنس', 'الحالة الاجتماعية', 'رقم الهاتف', 'العنوان',
               'الحالة', 'رقم هوية رب الأسرة (إن وجد)']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    # Add sample data for user guidance
    sample_data = ['أحمد', 'محمد', 'علي', 'الأحمد', '123456789', '1990-01-01',
                   'ذكر', 'متزوج', '0501234567', 'الرياض', 'مكتمل', '']
    sample_member_data = ['فاطمة', 'أحمد', 'محمد', 'الأحمد', '987654321', '2015-05-10',
                          'أنثى', 'أعزب', '', 'الرياض', 'مكتمل', '123456789']

    for col, data in enumerate(sample_data, 1):
        ws.cell(row=2, column=col, value=data)
    for col, data in enumerate(sample_member_data, 1):
        ws.cell(row=3, column=col, value=data)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='beneficiaries_template.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@bp.route('/export_records')
@login_required
def export_records():
    """
    Handles exporting beneficiary records to an Excel file.
    Can export all records or a filtered subset.
    """
    query = Record.query.filter_by(is_deleted=False)
    if request.args.get('type') == 'filtered':
        # Apply filters based on query parameters
        search = request.args.get('search', '')
        if search:
            search_term = f"%{search}%"
            query = query.filter(or_(
                Record.first_name.ilike(search_term), Record.father_name.ilike(search_term),
                Record.grandfather_name.ilike(search_term), Record.family_name.ilike(search_term),
                Record.id_passport_number.ilike(search_term), Record.phone_number.ilike(search_term)
            ))
        status_filter = request.args.get('status', '')
        if status_filter:
            query = query.filter_by(status=status_filter)
        start_date_str = request.args.get('start_date', '')
        if start_date_str:
            query = query.filter(Record.created_at >= datetime.strptime(start_date_str, '%Y-%m-%d').date())
        end_date_str = request.args.get('end_date', '')
        if end_date_str:
            query = query.filter(Record.created_at <= datetime.strptime(end_date_str, '%Y-%m-%d').date())

    records = query.order_by(Record.created_at.desc()).all()

    # --- Create Excel Workbook ---
    wb = Workbook()
    ws = wb.active
    ws.title = "Beneficiaries Export"
    headers = ['ID', 'الاسم الأول', 'اسم الأب', 'اسم الجد', 'اسم العائلة', 'رقم الهوية/جواز السفر',
               'تاريخ الميلاد', 'الجنس', 'الحالة الاجتماعية', 'رقم الهاتف', 'عدد أفراد الأسرة',
               'العنوان', 'الحالة', 'رقم هوية رب الأسرة', 'تاريخ الإنشاء']
    ws.append(headers)

    for record in records:
        head_of_household_passport = ''
        if record.head_of_household_id:
            head = Record.query.get(record.head_of_household_id)
            if head:
                head_of_household_passport = head.id_passport_number

        ws.append([
            record.id, record.first_name, record.father_name, record.grandfather_name,
            record.family_name, record.id_passport_number,
            record.date_of_birth.strftime('%Y-%m-%d'), record.gender, record.marital_status,
            record.phone_number or '', record.family_members_count, record.address or '',
            record.status, head_of_household_passport, record.created_at.strftime('%Y-%m-%d %H:%M')
        ])

    # --- Serve the file ---
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'beneficiaries_export_{timestamp}.xlsx'
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@bp.route('/generate_pdf/<int:record_id>')
@login_required
def generate_pdf(record_id):
    """
    Generates a professional PDF report for a single beneficiary.
    """
    record = Record.query.get_or_404(record_id)
    try:
        pdf_buffer = generate_beneficiary_pdf(record)
        return send_file(
            pdf_buffer,
            as_attachment=True,
            download_name=f'beneficiary_report_{record.id}.pdf',
            mimetype='application/pdf'
        )
    except Exception as e:
        flash(f'حدث خطأ أثناء إنشاء ملف PDF: {str(e)}', 'error')
        return redirect(url_for('routes.beneficiaries'))


# --- API Routes ---
# These routes are used for AJAX requests from the frontend to provide dynamic functionality.

@bp.route('/api/search_suggestions')
@login_required
def search_suggestions():
    """
    Provides live search suggestions for beneficiary names and IDs.
    """
    query = request.args.get('q', '').strip()
    if len(query) < 2: return jsonify([])
    search_term = f"%{query}%"
    records = Record.query.filter_by(is_deleted=False).filter(or_(
        Record.first_name.ilike(search_term), Record.father_name.ilike(search_term),
        Record.grandfather_name.ilike(search_term), Record.family_name.ilike(search_term),
        Record.id_passport_number.ilike(search_term), Record.phone_number.ilike(search_term)
    )).limit(5).all()
    suggestions = [{'name': r.full_name, 'id_passport': r.id_passport_number, 'phone': r.phone_number or '-'} for r in records]
    return jsonify(suggestions)

@bp.route('/api/statistics')
@login_required
def api_statistics():
    """
    Provides key statistics for the dashboard or other parts of the application.
    """
    return jsonify({
        'total': Record.query.filter_by(is_deleted=False).count(),
        'completed': Record.query.filter_by(status='مكتمل', is_deleted=False).count(),
        'review': Record.query.filter_by(status='بحاجة لمراجعة', is_deleted=False).count(),
        'inactive': Record.query.filter_by(status='غير نشط', is_deleted=False).count()
    })

@bp.route('/api/beneficiary/<int:record_id>')
@login_required
def api_get_beneficiary(record_id):
    """
    Returns the full data for a single beneficiary as JSON.
    Used to populate the 'Edit Beneficiary' modal.
    """
    record = Record.query.get_or_404(record_id)
    # Convert record to a dictionary, handling date serialization.
    record_data = {column.name: getattr(record, column.name) for column in record.__table__.columns}
    if isinstance(record_data.get('date_of_birth'), datetime):
        record_data['date_of_birth'] = record_data['date_of_birth'].strftime('%Y-%m-%d')
    elif record_data.get('date_of_birth') is not None:
        record_data['date_of_birth'] = record_data['date_of_birth'].isoformat()

    if isinstance(record_data.get('created_at'), datetime):
        record_data['created_at'] = record_data['created_at'].isoformat()
    if isinstance(record_data.get('updated_at'), datetime):
        record_data['updated_at'] = record_data['updated_at'].isoformat()

    record_data['is_deleted'] = record.is_deleted

    # Include a list of potential heads of household for the dropdown, excluding the current record and soft-deleted records.
    potential_heads = Record.query.filter(Record.id != record_id, Record.head_of_household_id.is_(None), Record.is_deleted == False).order_by(Record.first_name, Record.family_name).all()
    record_data['potential_heads_options'] = [
        {'id': p.id, 'full_name': p.full_name, 'id_passport_number': p.id_passport_number} for p in potential_heads
    ]
    return jsonify(record_data)

@bp.route('/api/beneficiary/<int:record_id>/family')
@login_required
def api_get_family_members(record_id):
    """
    Returns the family members for a given head of household as JSON.
    """
    head_of_household = Record.query.get_or_404(record_id)
    if head_of_household.head_of_household_id is not None:
        return jsonify({'error': 'Provided record is not a head of household'}), 400

    family_members = head_of_household.family_members.order_by(Record.date_of_birth).all()

    family_data = [
        {
            'id': member.id,
            'full_name': member.full_name,
            'id_passport_number': member.id_passport_number,
            'gender': member.gender,
            'age': member.age,
            'marital_status': member.marital_status,
            'status': member.status
        } for member in family_members
    ]
    return jsonify(family_data)

# --- Settings and User Management ---
@bp.route('/settings')
@admin_required
def settings():
    """
    Displays the user management page for admins.
    """
    users = User.query.order_by(User.created_at.desc()).all()
    return render_template('settings.html', users=users)

@bp.route('/add_user', methods=['POST'])
@admin_required
def add_user():
    """
    Handles the creation of a new user by an admin.
    """
    if User.query.filter_by(username=request.form['username']).first():
        flash('اسم المستخدم موجود بالفعل', 'error')
        return redirect(url_for('routes.settings'))
    try:
        user = User(username=request.form['username'], full_name=request.form['full_name'], role=request.form['role'])
        user.set_password(request.form['password'])
        db.session.add(user)
        db.session.commit()
        flash('تم إضافة المستخدم بنجاح', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء إضافة المستخدم: {str(e)}', 'error')
    return redirect(url_for('routes.settings'))

@bp.route('/edit_user/<int:user_id>', methods=['POST'])
@admin_required
def edit_user(user_id):
    """
    Handles editing an existing user's details by an admin.
    """
    user = User.query.get_or_404(user_id)
    # Prevent an admin from demoting themselves.
    if user.id == current_user.id and request.form['role'] != 'admin':
        flash('لا يمكنك تغيير دورك الإداري', 'error')
        return redirect(url_for('routes.settings'))
    try:
        if User.query.filter(User.username == request.form['username'], User.id != user_id).first():
            flash('اسم المستخدم موجود بالفعل', 'error')
            return redirect(url_for('routes.settings'))
        user.username = request.form['username']
        user.full_name = request.form['full_name']
        user.role = request.form['role']
        if request.form.get('password'):
            user.set_password(request.form['password'])
        db.session.commit()
        flash('تم تحديث بيانات المستخدم بنجاح', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء تحديث البيانات: {str(e)}', 'error')
    return redirect(url_for('routes.settings'))

@bp.route('/delete_user/<int:user_id>', methods=['POST'])
@admin_required
def delete_user(user_id):
    """
    Handles deleting a user by an admin.
    Prevents a user from deleting themselves or a user who has created records.
    """
    user = User.query.get_or_404(user_id)
    # Prevent self-deletion.
    if user.id == current_user.id:
        flash('لا يمكنك حذف حسابك الشخصي', 'error')
        return redirect(url_for('routes.settings'))
    # Prevent deletion of users with associated records to maintain data integrity.
    if Record.query.filter_by(created_by_user_id=user_id).first():
        flash('لا يمكن حذف المستخدم لأنه قام بإنشاء سجلات في النظام', 'error')
        return redirect(url_for('routes.settings'))
    try:
        db.session.delete(user)
        db.session.commit()
        flash('تم حذف المستخدم بنجاح', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'حدث خطأ أثناء حذف المستخدم: {str(e)}', 'error')
    return redirect(url_for('routes.settings'))